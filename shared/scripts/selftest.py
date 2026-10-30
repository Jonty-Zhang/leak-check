#!/usr/bin/env python3
"""自测：scan_leaks.py 抓得到该抓的、不误伤正常写法；fix_docx.py 改完格式还在、复扫干净。

    python3 shared/scripts/selftest.py
"""
from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

HERE = Path(__file__).resolve().parent
SAMPLES = HERE.parent / "tests" / "samples"
sys.path.insert(0, str(HERE))
import fix_docx  # noqa: E402
import scan_leaks  # noqa: E402

failed = 0


def check(cond, msg):
    global failed
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        failed += 1


def codes(doc, severity=None):
    out = set()
    for f in doc.findings:
        if severity is None or f.severity == severity:
            out.update(f.code.split("+"))
    return out


def errors_at(doc, needle):
    return [f for f in doc.findings if f.severity == "ERROR" and needle in f.excerpt]


def test_leaky_md():
    print("leaky.md：典型雷点都要报 ERROR")
    doc = scan_leaks.scan(str(SAMPLES / "leaky.md"))
    for needle in ("根据用户的要求", "用户说", "老师要求", "之后要替换掉这里的内容", "改进版"):
        check(errors_at(doc, needle), f"抓到“{needle}”")
    want = {"META_YOU", "CHAT_OPENER", "CHAT_CLOSER", "META_SELFCERT", "DRAFT_TODO", "DRAFT_PLACEHOLDER",
            "NUM_HEADING", "NUM_FIGREF", "CITE_RANGE", "REF_YEAR", "COUNT_MISMATCH"}
    missing = want - codes(doc)
    check(not missing, f"各类规则都触发（缺：{missing or '无'}）")


def test_clean_md():
    print("clean.md：正常论文里的“改进版算法”“待定系数法”“此处需要补充说明”“（第2版）”不误报")
    doc = scan_leaks.scan(str(SAMPLES / "clean.md"))
    check(not doc.findings, f"0 条（实际 {len(doc.findings)} 条：{[f.code for f in doc.findings]}）")


def test_user_domain():
    print("user-domain.md：需求分析里“用户希望……”是正常用法，不报 ERROR")
    doc = scan_leaks.scan(str(SAMPLES / "user-domain.md"))
    check(not codes(doc, "ERROR"), f"0 ERROR（实际 {codes(doc, 'ERROR')}）")
    check("META_USER" in codes(doc, "WARN"), "降为 WARN 提醒人工看")


def test_tex_and_stdin(tmp):
    print("tex / 标准输入")
    tex = tmp / "a.tex"
    tex.write_text("\\section{引言}\n% TODO 改一下\n根据用户的要求，本文研究 A。\n\\section{方法}\n", encoding="utf-8")
    doc = scan_leaks.scan(str(tex))
    check(any(f.code.startswith("DRAFT_TODO") and f.severity == "WARN" for f in doc.findings), "注释里的 TODO 降为 WARN")
    check(errors_at(doc, "根据用户的要求"), "tex 正文元话语报 ERROR")
    r = subprocess.run([sys.executable, str(HERE / "scan_leaks.py"), "-", "--json"], input="希望对你有帮助！\n",
                       capture_output=True, text=True)
    data = json.loads(r.stdout)
    check(r.returncode == 1 and data[0]["errors"] == 1, "stdin 能读，有 ERROR 时退出码为 1")


def build_docx(path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import parse_xml

    d = Document()
    d.add_heading("1 引言", level=1)
    p = d.add_paragraph()
    r1 = p.add_run("根据用户的")
    r1.bold = True
    r2 = p.add_run("要求，本文")
    r2.bold = True
    p.add_run("研究宿舍用电。")                              # 第 2 段：命中跨两个 run
    d.add_paragraph("最终版：")                               # 第 3 段
    p = d.add_paragraph("能耗公式为 $E=Pt$，详见")            # 第 4 段
    p.add_run("错误！未找到引用源。")
    p = d.add_paragraph("这一句被标黄了。")                   # 第 5 段
    p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    p = d.add_paragraph("已有修订：")                         # 第 6 段
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    p._p.append(parse_xml(f'<w:ins {w} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>新插入的字</w:t></w:r></w:ins>'))
    p._p.append(parse_xml(f'<w:del {w} w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>删掉的字</w:delText></w:r></w:del>'))
    p = d.add_paragraph("这里有一条批注。")                   # 第 7 段
    d.add_comment(p.runs, text="老师：这里再改改", author="T", initials="T")
    d.add_paragraph("图1 用电曲线", style="Caption")           # 第 8 段
    d.add_paragraph("图3 用电分布", style="Caption")           # 第 9 段
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "**加粗没转换**"                     # 第 10 段（表格内）
    d.add_paragraph("希望对你有帮助！")                       # 第 11 段
    d.sections[0].header.paragraphs[0].text = "宿舍用电调查（改进版）"
    d.save(str(path))


def test_docx(tmp):
    print("docx：扫描 + 修复 + 复扫")
    path = tmp / "report.docx"
    build_docx(path)
    doc = scan_leaks.scan(str(path))
    want = {"META_USER", "VERSION_MARK", "LATEX_RESIDUE", "BROKEN_FIELD", "DOCX_HIGHLIGHT", "DOCX_REVISION",
            "DOCX_COMMENT", "NUM_CAPTION", "MD_RESIDUE", "CHAT_CLOSER"}
    missing = want - codes(doc)
    check(not missing, f"docx 各类问题都抓到（缺：{missing or '无'}）")
    loc = {f.code.split("+")[0]: f.loc for f in doc.findings}
    check(loc.get("META_USER") == "第 2 段", f"段号正确（META_USER 在 {loc.get('META_USER')}）")
    check(any("页眉" in f.loc for f in doc.findings), "页眉里的版本标记也查到")
    check(any(f.loc.startswith("第 10 段") for f in doc.findings), "表格里的段落编号和正文连续")

    edits = tmp / "edits.json"
    edits.write_text(json.dumps([
        {"para": 2, "find": "根据用户的要求，", "replace": ""},
        {"para": 3, "delete": True, "find": "最终版"},
        {"para": 4, "find": "$E=Pt$", "replace": "E = Pt"},
        {"para": 4, "find": "，详见错误！未找到引用源。", "replace": "。"},
        {"para": 9, "find": "图3", "replace": "图2"},
        {"para": 10, "find": "**加粗没转换**", "replace": "加粗没转换"},
        {"para": 11, "delete": True},
        {"find": "（改进版）", "replace": ""},
    ], ensure_ascii=False), encoding="utf-8")

    bad = tmp / "bad_edits.json"
    bad.write_text(json.dumps([{"para": 2, "find": "不存在的句子", "replace": ""}], ensure_ascii=False), encoding="utf-8")
    before = path.read_bytes()
    with redirect_stdout(io.StringIO()):
        rc = fix_docx.main([str(path), str(bad), "--inplace"])
    check(rc == 1 and path.read_bytes() == before, "有一条没命中就整份不写入")

    with redirect_stdout(io.StringIO()):
        rc = fix_docx.main([str(path), str(edits), "--inplace", "--accept-revisions", "--remove-comments", "--clear-highlight"])
    check(rc == 0, "修复脚本执行成功")
    check((tmp / "report.leakfix-backup.docx").exists(), "原文件已备份")

    after = scan_leaks.scan(str(path))
    check(not codes(after, "ERROR"), f"复扫 0 ERROR（剩：{[(f.code, f.loc) for f in after.findings if f.severity == 'ERROR']}）")

    from docx import Document
    d = Document(str(path))
    p = d.paragraphs[1]
    check(p.text == "本文研究宿舍用电。", f"跨 run 替换结果正确（{p.text!r}）")
    check(p.runs[0].bold and p.runs[0].text.startswith("本文"), "被替换处保留了原来的加粗格式")
    check("新插入的字" in d.paragraphs[4].text and "删掉的字" not in d.element.xml, "修订已接受")
    check(not any(r.font.highlight_color for para in d.paragraphs for r in para.runs), "高亮已去掉")


def main():
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        test_leaky_md()
        test_clean_md()
        test_user_domain()
        test_tex_and_stdin(tmp)
        test_docx(tmp)
    print("\n全部通过" if not failed else f"\n{failed} 项失败")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
